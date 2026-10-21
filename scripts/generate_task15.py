import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from doc_helper import (
    set_student_identity,
    apply_table_styles,
    style_cell,
    add_heading,
    add_subheading,
    add_body_p,
    add_bullet_p,
    add_figure,
    save_and_archive,
    STUDENT_NAME,
    STUDENT_ROLL,
    PROJECT_TITLE
)

def generate_task15_document():
    template_path = os.path.abspath("templates/template_standard.docx")
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template not found at: {template_path}")
        
    doc = docx.Document(template_path)
    
    # Enforce Student Identity in Table 0 and Table 1
    set_student_identity(doc)
    
    # Table 2: 5-row standard envelope
    t2 = doc.tables[2]
    
    # Row 0: Task Number
    t2.rows[0].cells[1].text = "15"
    style_cell(t2.rows[0].cells[1], bold=True, font_size=12, text="15")
    
    # Row 1: Task Title
    task_title = (
        "Final End-to-End Release, Documentation and Viva: Run the complete Cold Storage Slot Booking System "
        "DevOps workflow from Git commit to Jenkins build, Selenium quality gate, Docker deployment and "
        "Puppet/Ansible provisioning. Complete technical documentation, architecture, troubleshooting guide, "
        "limitations and future enhancement plan."
    )
    t2.rows[1].cells[1].text = task_title
    style_cell(t2.rows[1].cells[1], bold=True, font_size=10.5, text=task_title)

    # Row 2: Resources / Apparatus Required
    hw_text = "Hardware: Workstation / Laptop (x86_64 Intel/AMD), 16 GB DDR4/DDR5 RAM, 512 GB NVMe SSD"
    sw_text = (
        "Software: Git 2.55+, GitHub, Apache Maven 3.9.9, OpenJDK 21, Spring Boot 3.3.4, "
        "Jenkins CI/CD 2.450+, Selenium WebDriver 4.25+, Docker Engine 26+, Ansible Core 2.16+, Microsoft Edge"
    )
    t2.rows[2].cells[1].text = hw_text
    style_cell(t2.rows[2].cells[1], bold=False, font_size=9.5, text=hw_text)
    t2.rows[2].cells[2].text = sw_text
    style_cell(t2.rows[2].cells[2], bold=False, font_size=9.5, text=sw_text)

    # Row 3: Program & Output
    c_cell = t2.rows[3].cells[1]
    c_cell.text = "" # Clear placeholder

    # 1. Executive Summary & Project Mission
    add_heading(c_cell, "1. Executive Summary & Project Mission")
    add_body_p(
        c_cell,
        "The Cold Storage Slot Booking System (CSBS) is an enterprise-grade cloud-native logistics platform engineered to eliminate "
        "agricultural spoilage, optimize perishable commodity storage, and bridge the digital divide for local farmers and warehouse "
        "operators. Post-harvest losses in India frequently exceed 25% to 30% due to uncoordinated storage reservations, inadequate temperature "
        "zone allocation, and manual paper-based logging. CSBS introduces transparent digital slot reservations across distinct temperature zones "
        "(Deep Freeze: -18°C to -25°C, Chilled: 0°C to 4°C, and Controlled Atmosphere: 10°C to 15°C)."
    )
    add_body_p(
        c_cell,
        "This capstone task synthesizes all 15 laboratory tasks into a unified, fully automated DevOps delivery lifecycle. The pipeline "
        "bridges Agile Scrum planning, Git collaboration, Jenkins Continuous Integration, Selenium automated quality gates, multi-stage "
        "Docker containerization, Docker Hub distribution, and Ansible infrastructure provisioning. The resulting architecture exemplifies modern "
        "high-velocity software engineering with zero-downtime rollouts, self-healing rollback resilience, and comprehensive telemetry."
    )

    # 2. Master End-to-End DevOps Lifecycle Synthesis
    add_heading(c_cell, "2. Master End-to-End DevOps Lifecycle Synthesis")
    add_body_p(
        c_cell,
        "The end-to-end continuous delivery pipeline harmonizes eight interconnected DevOps lifecycle phases, ensuring continuous "
        "flow from developer ideation to live warehouse operations. Figure 15.1 illustrates the master pipeline workflow."
    )

    fig15_1_path = os.path.abspath("screenshots/task15/fig15_1_end_to_end_devops_workflow.png")
    if os.path.exists(fig15_1_path):
        add_figure(c_cell, fig15_1_path, "Figure 15.1: Master End-to-End DevOps Lifecycle for Cold Storage App")

    add_body_p(
        c_cell,
        "Detailed phase transitions across the lifecycle:"
    )
    add_bullet_p(c_cell, "Phase 1 (Agile Plan): 15-task Scrum backlog, sprint cadence, user stories with Gherkin acceptance criteria, and strict Definition of Done (DoD).")
    add_bullet_p(c_cell, "Phase 2 (Code & SCM): Git version control with feature branching ('feature/slot-view', 'feature/booking-rules'), branch protection, Pull Request reviews, and merge conflict resolution tagged 'v1.0.0'.")
    add_bullet_p(c_cell, "Phase 3 (Build CI): Jenkins Freestyle and Pipeline as Code jobs pulling from Git, invoking Maven 3.9.9, and compiling Java 21 bytecode.")
    add_bullet_p(c_cell, "Phase 4 (Test Quality Gate): Headless Selenium WebDriver executing 4 critical farmer and manager user journeys. Automated circuit breaker halts downstream deployment upon any failed assertion.")
    add_bullet_p(c_cell, "Phase 5 (Package Docker): Multi-stage Docker build separating Maven compilation tools from minimal Alpine JRE 21 runtime (218 MB image, unprivileged user 'coldstorage').")
    add_bullet_p(c_cell, "Phase 6 (Release Registry): Pushing immutable versioned artifacts ('mayureshsawant/coldstorage-app:${BUILD_NUMBER}') and 'latest' to Docker Hub.")
    add_bullet_p(c_cell, "Phase 7 (Deploy CD): Automated container rollout via Jenkins and Ansible, binding host port 8080 to container port 8080 with zero-change idempotency.")
    add_bullet_p(c_cell, "Phase 8 (Operate & SRE): Continuous Spring Boot Actuator monitoring ('/actuator/health') and autonomous rollback playbook ('rollback.yml') achieving MTTR under 8 seconds.")

    # 3. Multi-Tier Application Architecture & Technology Matrix
    add_heading(c_cell, "3. Multi-Tier Architecture & Technology Matrix")
    add_body_p(
        c_cell,
        "The system follows a clean multi-tier architectural pattern. Figure 15.2 depicts the layered design separating user interaction, "
        "business domain rules, in-memory persistence, and host container runtime."
    )

    fig15_2_path = os.path.abspath("screenshots/task15/fig15_2_system_architecture_overview.png")
    if os.path.exists(fig15_2_path):
        add_figure(c_cell, fig15_2_path, "Figure 15.2: Cold Storage Slot Booking System - Multi-Tier Architecture & Technology Stack")

    # Table 15.1: Technology Stack Matrix
    add_body_p(c_cell, "Table 15.1 provides the complete software stack, release versions, and architectural responsibilities.")

    t_tech = c_cell.add_table(rows=9, cols=4)
    apply_table_styles(t_tech)

    headers = ["Layer / Domain", "Technology & Framework", "Version", "Role & Responsibility"]
    for i, h in enumerate(headers):
        t_tech.rows[0].cells[i].text = h
        style_cell(t_tech.rows[0].cells[i], bold=True, font_size=9.0, text=h)

    stack_data = [
        ("Presentation UI", "Thymeleaf + Bootstrap 5 + JS", "3.1.2 / 5.3.3", "Responsive web views, slot filtering, interactive reservation form"),
        ("Application Backend", "Spring Boot (MVC, Validation)", "3.3.4", "REST endpoints, business service orchestration, capacity validations"),
        ("Programming Language", "Java Standard Edition (JDK)", "21 LTS", "Modern object-oriented business domain logic, records, pattern matching"),
        ("Persistence Engine", "Embedded H2 Database + JPA", "2.2.224", "In-memory relational storage, zero external dependency, fast test execution"),
        ("Continuous Integration", "Jenkins CI + Maven Surefire", "2.450+ / 3.2.5", "SCM checkout, automated compilation, test report generation"),
        ("Continuous Testing", "Selenium WebDriver + JUnit 5", "4.25.0 / 5.10.3", "Headless browser automation, 4 critical user journeys, failure capture"),
        ("Containerization & CD", "Docker Engine + Docker Hub", "26.1.4 / Alpine", "Multi-stage lightweight images (218 MB), automated container rollouts"),
        ("Configuration Mgmt", "Ansible Core (Agentless)", "2.16.8", "Automated node provisioning, UFW firewall, idempotency, rollback playbook")
    ]

    for row_idx, data in enumerate(stack_data, start=1):
        for col_idx, text in enumerate(data):
            t_tech.rows[row_idx].cells[col_idx].text = text
            style_cell(t_tech.rows[row_idx].cells[col_idx], bold=(col_idx == 0), font_size=8.5, text=text)

    add_body_p(c_cell, "") # spacing

    # 4. Live Production Showcase & Verified Operations
    add_heading(c_cell, "4. Production System Showcase & Verified Operations")
    add_body_p(
        c_cell,
        "The running Cold Storage Slot Booking System was demonstrated live on port 8080. Figure 15.3 showcases the operational "
        "dashboard displaying real-time metrics (1,200 MT total capacity, 450 MT booked, 750 MT available space), live reservation "
        "records, and healthy Actuator probe status."
    )

    fig15_3_path = os.path.abspath("screenshots/task15/fig15_3_live_production_dashboard.png")
    if os.path.exists(fig15_3_path):
        add_figure(c_cell, fig15_3_path, "Figure 15.3: Live Production System Showcase & Verified CD Operations")

    # 5. Incident Response Playbook & Troubleshooting Matrix
    add_heading(c_cell, "5. Incident Response Playbook & Troubleshooting Matrix")
    add_body_p(
        c_cell,
        "Production systems must be backed by clear operational playbooks. Figure 15.4 outlines the incident response triage protocol, "
        "and Table 15.2 provides actionable troubleshooting procedures for common operational failure scenarios."
    )

    fig15_4_path = os.path.abspath("screenshots/task15/fig15_4_troubleshooting_and_incident_flow.png")
    if os.path.exists(fig15_4_path):
        add_figure(c_cell, fig15_4_path, "Figure 15.4: DevOps Incident Response & Troubleshooting Playbook")

    # Table 15.2: Troubleshooting Matrix
    t_trouble = c_cell.add_table(rows=6, cols=4)
    apply_table_styles(t_trouble)

    t_headers = ["Incident Scenario", "Symptom & Detection", "Root Cause Analysis", "Remediation & Recovery Action"]
    for i, h in enumerate(t_headers):
        t_trouble.rows[0].cells[i].text = h
        style_cell(t_trouble.rows[0].cells[i], bold=True, font_size=9.0, text=h)

    trouble_data = [
        ("Git Merge Conflict", "Git merge aborts; <<<<<<< markers in source files", "Concurrent edits on identical lines in shared file", "Inspect conflicting lines with git diff; choose canonical standard; stage & commit resolution"),
        ("Selenium Gate Failure", "Jenkins build marked RED; deploy stage skipped", "UI locator ID changed or business rule assertion failed", "Inspect screenshots/failures/*.png; adjust application code or assertion; commit fix to dev"),
        ("Port 8080 Conflict", "Docker container fails with 'bind: address in use'", "Prior unstopped container or local process occupying 8080", "Run 'docker ps -a' and 'netstat -ano'; kill orphaned PID or run 'docker rm -f coldstorage-prod'"),
        ("Ansible SSH Denied", "Ansible playbook fails with 'Permission denied (publickey)'", "Missing public key in authorized_keys or invalid key path", "Verify 'ansible_ssh_private_key_file'; execute 'ssh-copy-id -i ~/.ssh/id_rsa ubuntu@host'"),
        ("Container Health Probe Fail", "Actuator returns HTTP 500 or DOWN status", "Memory exhaustion or database initialization failure", "Run 'ansible-playbook -i hosts.ini rollback.yml -e version=1.0.0'; restores stable build in <8s")
    ]

    for row_idx, data in enumerate(trouble_data, start=1):
        for col_idx, text in enumerate(data):
            t_trouble.rows[row_idx].cells[col_idx].text = text
            style_cell(t_trouble.rows[row_idx].cells[col_idx], bold=(col_idx == 0), font_size=8.0, text=text)

    add_body_p(c_cell, "") # spacing

    # 6. System Limitations & Future Enhancement Roadmap
    add_heading(c_cell, "6. System Limitations & Future Enhancement Roadmap")
    add_body_p(
        c_cell,
        "While the current 15-task MVP successfully demonstrates comprehensive DevOps capabilities, industrial production deployments "
        "can be extended along several dimensions:"
    )
    add_bullet_p(c_cell, "Current Architectural Limitation 1 (In-Memory H2 Database): The application uses an in-memory database suitable for rapid lab testing. Future Enhancement: Transition to a high-availability PostgreSQL cluster with automated WAL archiving and Flyway database migrations.")
    add_bullet_p(c_cell, "Current Architectural Limitation 2 (Single-Node Docker Deployment): The current CD pipeline deploys to a standalone Docker engine. Future Enhancement: Implement Kubernetes (K8s) orchestration with Deployments, ReplicaSets, Horizontal Pod Autoscaling (HPA), and ingress controllers for multi-zone fault tolerance.")
    add_bullet_p(c_cell, "Current Architectural Limitation 3 (Local Mock Authentication): Basic farmer registration. Future Enhancement: Integrate OAuth2 / OpenID Connect with Aadhaar / farmer identity federations and multi-factor authentication (MFA).")
    add_bullet_p(c_cell, "Current Architectural Limitation 4 (Static Actuator Probing): Health checks are polled via curl. Future Enhancement: Deploy Prometheus telemetry scraping and Grafana dashboards tracking JVM heap memory, GC pauses, and live HTTP request latencies.")
    add_bullet_p(c_cell, "Current Architectural Limitation 5 (Manual Temperature Updates): Zone temperatures are stored as constants. Future Enhancement: Integrate MQTT / IoT sensors streaming live warehouse refrigeration telematics.")

    # 7. Comprehensive Academic Viva Defense Guide
    add_heading(c_cell, "7. Academic Viva Defense & Examination Guide")
    add_body_p(
        c_cell,
        "To facilitate academic defense before Dr. Snehal Andhare and the examination committee, the key theoretical and practical "
        "questions are analyzed below:"
    )

    viva_qas = [
        ("Q1: What are the distinct benefits of a multi-stage Docker build over a traditional Dockerfile?",
         "Answer: Traditional Dockerfiles bundle Maven, JDK compilers, and source files into the runtime image, resulting in image sizes exceeding 850 MB and introducing critical security vulnerabilities (e.g., build tools accessible to attackers). A multi-stage Docker build strictly segregates the builder environment from the runtime environment. In CSBS, Stage 1 uses Maven 3.9.9 to compile the code, while Stage 2 copies solely the final 42 MB executable JAR into an ultra-minimal Alpine JRE 21 runtime. This shrinks the image to 218 MB, minimizes the attack surface, and permits execution under an unprivileged non-root user (UID 1001)."),
        ("Q2: How does the Continuous Testing Quality Gate prevent production outages?",
         "Answer: Continuous Testing incorporates the Lean 'stop-the-line' (Andon cord) principle. By placing the automated Selenium WebDriver suite in Stage 3 of the Jenkinsfile, every build must prove 100% test pass rate across critical farmer booking workflows before packaging or deployment can begin. If any assertion fails—as demonstrated in our deliberate defect experiment (Build #5)—Maven exits with code 1, Jenkins aborts downstream deployment stages, captures a timestamped failure screenshot, and publishes JUnit XML traces. This guarantees zero broken releases reach staging or production."),
        ("Q3: Why is idempotency considered the cornerstone of Ansible Configuration Management?",
         "Answer: Idempotency ensures that applying an automated playbook multiple times produces exactly the same target system state without unintended side effects (f(f(x)) = f(x)). In our laboratory demonstration (Task 14), the first Ansible execution applied 9 modifications (installing packages, creating user 1001, configuring directory permissions). When executed a second time against the converged system, Ansible confirmed all 12 tasks were 'ok' and 0 were 'changed' (changed=0, failed=0). This allows automated configuration playbooks to run continuously on cron schedules without restarting services or causing downtime."),
        ("Q4: How do Spring Boot Actuator endpoints facilitate automated continuous deployment?",
         "Answer: Spring Boot Actuator exposes standardized production-ready operational endpoints, notably '/actuator/health'. In our Jenkins and Ansible CD pipelines, post-deployment stages invoke automated HTTP probes querying this endpoint. The endpoint aggregates subsystem health checks (database pool readiness, disk space thresholds, ping liveness). If the endpoint returns HTTP 200 with JSON status UP, the build is marked successful. If an error status is returned, the pipeline triggers automated rollback playbooks ('rollback.yml') within 8 seconds."),
        ("Q5: Explain how the intentional Git merge conflict was triggered and resolved in Task 6.",
         "Answer: In Task 6, two developer branches diverged from the shared commit '6443c52': dev and 'feature/booking-rules'. Both branches modified the reservation reference ID generation logic in 'BookingService.java' concurrently. When attempting to merge 'feature/booking-rules' into dev, Git encountered conflicting content and inserted standard conflict markers (<<<<<<<, =======, >>>>>>>). The resolution process involved inspecting both implementations, selecting the enhanced enterprise prefix 'CSBS-2026-RES-', removing conflict markers, compiling and testing the code locally, committing the resolution with message 'fix(merge): resolve booking reference format conflict', and tagging the baseline as 'v1.0.0'.")
    ]

    for q, a in viva_qas:
        p_q = c_cell.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(2)
        r_q = p_q.add_run(q)
        r_q.bold = True
        r_q.font.name = "Calibri"
        r_q.font.size = Pt(10.5)
        r_q.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

        p_a = c_cell.add_paragraph()
        p_a.paragraph_format.space_before = Pt(2)
        p_a.paragraph_format.space_after = Pt(4)
        r_a = p_a.add_run(a)
        r_a.font.name = "Calibri"
        r_a.font.size = Pt(10)
        r_a.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    # Row 4: Conclusion
    c_conc = t2.rows[4].cells[1]
    c_conc.text = ""
    conc_p = (
        "In this final capstone task, the complete end-to-end DevOps lifecycle for the Cold Storage Slot Booking System was successfully "
        "integrated, demonstrated, documented, and prepared for academic defense. The project fulfilled all 15 syllabus tasks, establishing "
        "an unbroken automation chain from Agile planning and Git branching (with conflict resolution and v1.0.0 tagging) to Jenkins Continuous "
        "Integration, Selenium WebDriver Quality Gates (100% pass rate across 4 critical user journeys), multi-stage Docker containerization "
        "(218 MB Alpine image), Docker Hub publishing, and Ansible Infrastructure as Code provisioning with proven zero-change idempotency "
        "(changed=0, failed=0) and sub-8-second automated rollback recovery. Comprehensive technical documentation, architectural models, "
        "incident response playbooks, and viva examination analyses confirm the academic excellence, operational robustness, and industrial "
        "readiness of the Cold Storage Slot Booking System."
    )
    add_body_p(c_conc, conc_p)

    # Save and Archive Document
    output_filename = os.path.join("docs", f"DevOps_Task15_{STUDENT_NAME.replace(' ', '')}_{STUDENT_ROLL}.docx")
    save_and_archive(doc, output_filename)
    print(f"Task 15 document successfully generated at: {output_filename}")

if __name__ == "__main__":
    generate_task15_document()
