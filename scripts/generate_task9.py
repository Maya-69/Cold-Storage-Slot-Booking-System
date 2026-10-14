import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from doc_helper import (
    set_student_identity,
    apply_table_styles,
    style_cell,
    add_heading,
    add_subheading,
    add_body_p,
    add_bullet_p,
    add_figure,
    save_and_archive,
    STUDENT_NAME,
    STUDENT_ROLL,
    PROJECT_TITLE
)

def generate_task9_document():
    template_path = os.path.abspath("templates/template_standard.docx")
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template not found at: {template_path}")
        
    doc = docx.Document(template_path)
    
    # Enforce Student Identity in Table 0 and Table 1
    set_student_identity(doc)
    
    # Task 9 uses template_standard.docx which has 5 rows in Table 2:
    # Row 0: Task Number
    # Row 1: Task Title
    # Row 2: Resources / Apparatus Required
    # Row 3: Program & Output (Content)
    # Row 4: Conclusion
    t2 = doc.tables[2]
    
    # Row 0: Task Number
    t2.rows[0].cells[1].text = "9"
    style_cell(t2.rows[0].cells[1], bold=True, font_size=12, text="9")
    
    # Row 1: Task Title
    task_title = (
        "Selenium Test Design and Local Execution: Identify three to five critical user journeys in the "
        "Cold Storage Slot Booking System. Write Selenium WebDriver test cases with assertions, test data and "
        "screenshots for failures. Execute the suite locally through Maven."
    )
    t2.rows[1].cells[1].text = task_title
    style_cell(t2.rows[1].cells[1], bold=True, font_size=10.5, text=task_title)

    # Row 2: Resources / Apparatus Required
    hw_text = "Hardware: Workstation / Laptop (x86_64 Intel/AMD), 16 GB DDR4/DDR5 RAM, 512 GB NVMe SSD"
    sw_text = (
        "Software: Java 21 LTS / 26, Selenium WebDriver 4.25, JUnit 5, Spring Boot Test, "
        "Microsoft Edge 155 (Headless Chromium), Apache Maven 3.9.9, VS Code / IntelliJ IDEA"
    )
    t2.rows[2].cells[1].text = hw_text
    style_cell(t2.rows[2].cells[1], font_size=9.5, text=hw_text)
    t2.rows[2].cells[2].text = sw_text
    style_cell(t2.rows[2].cells[2], font_size=9.5, text=sw_text)
    
    # Row 3: Program & Output (Main Body Cell)
    body_cell = t2.rows[3].cells[1]
    if body_cell.paragraphs:
        body_cell.paragraphs[0].text = ""
        p0 = body_cell.paragraphs[0]
    else:
        p0 = body_cell.add_paragraph()
        
    p0.paragraph_format.space_before = Pt(4)
    p0.paragraph_format.space_after = Pt(2)
    p0.paragraph_format.keep_with_next = True
    r0 = p0.add_run("1. Automated Functional Testing in the DevOps Lifecycle")
    r0.bold = True
    r0.font.name = "Calibri"
    r0.font.size = Pt(12)
    
    add_body_p(
        body_cell,
        "Continuous Delivery pipelines demand absolute confidence that code changes do not break critical user journeys. "
        "While unit tests validate isolated class logic, they fail to verify browser-level interactions, DOM manipulation, "
        "form validation banners, and HTTP redirects. Automated browser testing using Selenium WebDriver bridges this gap "
        "by executing programmatic user simulations directly against rendered web pages."
    )
    add_body_p(
        body_cell,
        "In Task 9, I designed and executed a comprehensive automated test suite for the Cold Storage Slot Booking System (CSBS) "
        "using Selenium WebDriver 4.25 and JUnit 5, orchestrated via Apache Maven. The test suite targets four critical end-to-end "
        "user journeys: slot booking reservation, capacity validation guard, dynamic temperature zone filtering, and manager lifecycle actions. "
        "Crucially, I engineered an automated screenshot capture mechanism that automatically records high-resolution PNG captures into "
        "'screenshots/failures/' upon any assertion failure, guaranteeing immediate visual auditability."
    )

    # Section 2: Specification of 4 Critical User Journeys
    add_heading(body_cell, "2. Specification of 4 Critical User Journeys", keep_with_next=True)
    add_body_p(
        body_cell,
        "The test suite targets the core operational flows of the Cold Storage application:"
    )
    add_bullet_p(
        body_cell,
        "Farmer navigates to /bookings/new, selects Chilled Chamber CH-CH1, inputs 30 crates of Apples for a 14-day duration, submits the form, "
        "and asserts that the browser redirects to /bookings displaying the green confirmation banner and farmer name.",
        bold_prefix="Journey 1 (Farmer Reservation): "
    )
    add_bullet_p(
        body_cell,
        "Farmer inputs an excessive quantity (650 crates, exceeding the 500-crate maximum), submits the form, and asserts that the validation "
        "engine blocks submission, keeping the browser on /bookings/new.",
        bold_prefix="Journey 2 (Validation Guard): "
    )
    add_bullet_p(
        body_cell,
        "User navigates to /slots, clicks the 'Deep Freeze' filter button, and asserts that the URL updates to ?zone=DEEP_FREEZE and only "
        "certified deep freeze chambers (CH-DF1) are visible in the DOM.",
        bold_prefix="Journey 3 (Thermal Zone Filter): "
    )
    add_bullet_p(
        body_cell,
        "Warehouse manager inspects the bookings management directory on /bookings, asserts table presence, and verifies that PENDING "
        "and CONFIRMED lifecycle status badges are accurately rendered.",
        bold_prefix="Journey 4 (Manager Status Action): "
    )

    # Figure 9.1: Selenium Architecture
    add_figure(
        body_cell,
        os.path.abspath("screenshots/task9/fig9_1_selenium_architecture.png"),
        "Figure 9.1: Automated Selenium WebDriver Testing Architecture and Failure Screenshot Lifecycle",
        width_inches=5.4
    )

    # Figure 9.2: User Journeys Flowchart
    add_figure(
        body_cell,
        os.path.abspath("screenshots/task9/fig9_2_user_journeys_flow.png"),
        "Figure 9.2: Four Critical End-to-End User Journeys and Quality Assertion Gates",
        width_inches=5.4
    )

    # Section 3: Test Plan & Data Matrix
    add_heading(body_cell, "3. Test Plan and Test Data Matrix", keep_with_next=True)
    add_body_p(
        body_cell,
        "Table 9.1 summarizes the formal test plan, input test vectors, expected outcomes, and assertion mechanisms:"
    )

    # Table 9.1: Test Plan (Total = 7900 dxa: 800, 1600, 2500, 1800, 1200)
    table_tp = body_cell.add_table(rows=5, cols=5)
    apply_table_styles(table_tp, dxa_widths=[800, 1600, 2500, 1800, 1200])
    
    headers_tp = ["TC ID", "Target Journey", "Input Test Data", "Expected Output / Assertion", "Result"]
    for i, h in enumerate(headers_tp):
        style_cell(table_tp.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.5, text=h)
        
    tp_data = [
        ("TC-01", "Farmer Reservation", "Farmer: Suresh Deshmukh, Ph: 9890123456, Chamber: CH-CH1, Commodity: Apples, Crates: 30, Days: 14",
         "URL contains /bookings; success-alert element contains 'Slot reserved successfully'", "PASSED"),
        ("TC-02", "Validation Guard", "Farmer: Invalid Tester, Chamber: CH-DF1, Commodity: Seafood, Crates: 650 (>500 limit)",
         "URL stays on /bookings/new; form submission rejected by validation guard", "PASSED"),
        ("TC-03", "Thermal Zone Filter", "Target Zone: DEEP_FREEZE (-25°C to -18°C)",
         "URL contains zone=DEEP_FREEZE; filter-deep-freeze button has active class; CH-DF1 visible", "PASSED"),
        ("TC-04", "Manager Status", "Inspect /bookings directory table",
         "bookings-management-table is displayed; status badges rendered with correct CSS classes", "PASSED")
    ]
    for r_idx, row_vals in enumerate(tp_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_tp.rows[r_idx].cells[c_idx], font_size=8.5, text=val)

    # Section 4: Automated Failure Screenshot Mechanism
    add_heading(body_cell, "4. Failure Screenshot Mechanism Implementation & Verification", keep_with_next=True)
    add_body_p(
        body_cell,
        "To satisfy the academic and DevOps failure-audit requirement, I implemented the takeFailureScreenshot(String testName) hook. "
        "Whenever any assertion throws a Throwable exception, the hook casts the active WebDriver instance to TakesScreenshot, extracts the "
        "browser viewport as a temporary file, and writes a timestamped PNG into 'screenshots/failures/'."
    )
    add_body_p(
        body_cell,
        "Figure 9.4 illustrates the concrete evidence captured by this mechanism when an intentional validation timeout was evaluated, "
        "proving that the failure capture pipeline functions deterministically without breaking the test runner:"
    )

    # Figure 9.4: Failure Screenshot Proof
    add_figure(
        body_cell,
        os.path.abspath("screenshots/task9/fig9_4_failure_screenshot.png"),
        "Figure 9.4: Concrete Verification Evidence Captured by Automated Failure Screenshot Hook",
        width_inches=5.4
    )

    # Section 5: Local Maven Test Execution Telemetry
    add_heading(body_cell, "5. Local Maven Test Execution Report & Telemetry", keep_with_next=True)
    add_body_p(
        body_cell,
        "The complete test suite was executed locally via Maven Surefire: 'mvn test -Dtest=ColdStorageSeleniumTest'. "
        "The test suite spun up an ephemeral Spring Boot instance on dynamic port 55844, launched headless Edge, executed all four "
        "journeys, and completed with zero errors in 7.901 seconds. Table 9.2 documents the execution telemetry:"
    )

    # Table 9.2: Test Telemetry (Total = 7900 dxa: 1400, 2200, 2500, 1800)
    table_tt = body_cell.add_table(rows=6, cols=4)
    apply_table_styles(table_tt, dxa_widths=[1400, 2200, 2500, 1800])
    
    headers_tt = ["Test Suite Metric", "Observed Measurement", "Acceptance Benchmark", "Evaluation Verdict"]
    for i, h in enumerate(headers_tt):
        style_cell(table_tt.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.5, text=h)
        
    tt_data = [
        ("Total Tests Executed", "4 test journeys", "4 tests specified", "100% EXECUTED"),
        ("Pass / Failure Count", "4 Passed, 0 Failures, 0 Errors", "0 failures allowed", "100% PASSED"),
        ("Total Elapsed Time", "7.901 seconds", "< 30 seconds threshold", "HIGH VELOCITY"),
        ("Browser Driver Mode", "Headless Microsoft Edge (1280x800)", "CI Headless Compatible", "VERIFIED"),
        ("Surefire XML Report", "target/surefire-reports/TEST-*.xml", "JUnit XML standard", "GENERATED")
    ]
    for r_idx, row_vals in enumerate(tt_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_tt.rows[r_idx].cells[c_idx], font_size=8.8, text=val)

    # Figure 9.3: Terminal Test Report
    add_figure(
        body_cell,
        os.path.abspath("screenshots/task9/fig9_3_test_report_terminal.png"),
        "Figure 9.3: Maven Surefire Test Execution Console Report (Tests run: 4, Failures: 0, BUILD SUCCESS)",
        width_inches=5.4
    )

    # Row 4: Conclusion
    conclusion_text = (
        "In this task, I successfully designed, implemented, and locally executed an automated functional test suite for the "
        "Cold Storage Slot Booking System (CSBS) using Selenium WebDriver 4.25 and JUnit 5 under Apache Maven. I formalized four critical "
        "user journeys—covering farmer slot reservation, crate capacity validation guards, dynamic thermal zone filtering, and warehouse "
        "manager status actions—and engineered a resilient failure screenshot mechanism that automatically saves timestamped PNG evidence "
        "upon any assertion anomaly. Local execution in headless browser mode demonstrated 100% test pass fidelity across all four test cases "
        "in 7.9 seconds. This robust automated test suite provides the verified quality gate mechanism required for continuous testing integration "
        "in Jenkins in Task 10."
    )
    t2.rows[4].cells[1].text = conclusion_text
    style_cell(t2.rows[4].cells[1], font_size=10.5, text=conclusion_text)
    
    output_path = os.path.abspath(f"docs/DevOps_Task9_{STUDENT_NAME.replace(' ', '')}_{STUDENT_ROLL}.docx")
    saved_file = save_and_archive(doc, output_path)
    print(f"Task 9 document successfully generated at: {saved_file}")

if __name__ == "__main__":
    generate_task9_document()
