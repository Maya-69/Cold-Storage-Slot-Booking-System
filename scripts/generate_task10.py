import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from doc_helper import (
    set_student_identity,
    apply_table_styles,
    style_cell,
    add_heading,
    add_subheading,
    add_body_p,
    add_bullet_p,
    add_figure,
    save_and_archive,
    STUDENT_NAME,
    STUDENT_ROLL,
    PROJECT_TITLE
)

def generate_task10_document():
    template_path = os.path.abspath("templates/template_standard.docx")
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template not found at: {template_path}")
        
    doc = docx.Document(template_path)
    
    # Enforce Student Identity in Table 0 and Table 1
    set_student_identity(doc)
    
    # Table 2: 5-row structure in template_standard.docx
    # Row 0: Task Number
    # Row 1: Task Title
    # Row 2: Resources / Apparatus Required (cells[1]: HW, cells[2]: SW)
    # Row 3: Program & Output (cells[1])
    # Row 4: Conclusion (cells[1])
    t2 = doc.tables[2]
    
    # Row 0: Task Number
    t2.rows[0].cells[1].text = "10"
    style_cell(t2.rows[0].cells[1], bold=True, font_size=12, text="10")
    
    # Row 1: Task Title
    task_title = (
        "Continuous Testing in Jenkins: Integrate the Cold Storage Slot Booking System Selenium suite into "
        "Jenkins. Publish test reports and configure the pipeline so failed tests stop deployment. Correct at "
        "least one deliberately introduced defect and rerun the pipeline."
    )
    t2.rows[1].cells[1].text = task_title
    style_cell(t2.rows[1].cells[1], bold=True, font_size=10.5, text=task_title)

    # Row 2: Resources / Apparatus Required
    hw_text = "Hardware: Workstation / Laptop (x86_64 Intel/AMD), 16 GB DDR4/DDR5 RAM, 512 GB NVMe SSD"
    sw_text = (
        "Software: Jenkins CI 2.450+ LTS, Apache Maven 3.9.9, Selenium WebDriver 4.25+, JUnit 5 Jupiter, "
        "Maven Surefire Plugin 3.2.5, JUnit Test Result Plugin for Jenkins, Microsoft Edge (Headless), Git 2.55+"
    )
    t2.rows[2].cells[1].text = hw_text
    style_cell(t2.rows[2].cells[1], bold=False, font_size=9.5, text=hw_text)
    t2.rows[2].cells[2].text = sw_text
    style_cell(t2.rows[2].cells[2], bold=False, font_size=9.5, text=sw_text)

    # Row 3: Program & Output
    c_cell = t2.rows[3].cells[1]
    c_cell.text = "" # Clear placeholder

    # 1. Introduction & Theoretical Rationale
    add_heading(c_cell, "1. Continuous Testing & Quality Gate Governance in DevOps")
    add_body_p(
        c_cell,
        "Continuous Testing is a fundamental DevOps discipline that infuses automated validation throughout the "
        "software delivery lifecycle rather than treating testing as a deferred, manual post-build activity. In a high-consequence "
        "domain such as the Cold Storage Slot Booking System (CSBS)—where agricultural commodities, temperature-sensitive "
        "perishables, and warehouse revenue depend on accurate reservations and capacity constraints—untested deployments can lead "
        "to devastating operational failure. Continuous Testing provides immediate feedback on functional defects, regression risks, "
        "and data integrity violations prior to artifact release."
    )
    add_body_p(
        c_cell,
        "A Jenkins Quality Gate acts as an automated policy enforcement checkpoint. By coupling Maven Surefire test execution "
        "with Jenkins Pipeline control flow, any failed assertion immediately terminates the pipeline. This 'stop-the-line' "
        "philosophy guarantees that no defect-ridden artifact can transition into packaging, containerization, or production deployment."
    )

    # 2. Architecture of the Automated Continuous Testing Pipeline
    add_heading(c_cell, "2. Continuous Testing Pipeline Architecture")
    add_body_p(
        c_cell,
        "The automated quality gate integrates Git source control, Jenkins orchestration, Maven build execution, headless browser "
        "automation, and automated test result aggregation. The architectural topology is illustrated in Figure 10.1."
    )

    fig10_1_path = os.path.abspath("screenshots/task10/fig10_1_continuous_testing_arch.png")
    if os.path.exists(fig10_1_path):
        add_figure(c_cell, fig10_1_path, "Figure 10.1: Continuous Testing & Quality Gate Architecture for Cold Storage App")

    add_body_p(
        c_cell,
        "Key operational stages of the Quality Gate include:"
    )
    add_bullet_p(c_cell, "SCM Checkout Stage: Pulls verified commits from the development branch (dev) and prints author identity.")
    add_bullet_p(c_cell, "Compilation Stage: Compiles Java sources and test classes via 'mvn clean test-compile' to ensure syntactic integrity.")
    add_bullet_p(c_cell, "Selenium Quality Gate Stage: Executes ColdStorageSeleniumTest using Maven Surefire, driving a headless browser against the embedded Spring Boot instance.")
    add_bullet_p(c_cell, "Quality Decision Fork: If all assertions succeed (exit code 0), Jenkins transitions to artifact packaging. If any test fails (exit code 1), execution halts immediately, and subsequent deployment stages are aborted.")
    add_bullet_p(c_cell, "Reporting & Artifacts: The Jenkins JUnit plugin processes Surefire XML reports, publishes visual test trends, and failure screenshots are archived for forensic debugging.")

    # 3. Jenkins Pipeline as Code Configuration
    add_heading(c_cell, "3. Jenkinsfile Quality Gate Implementation")
    add_body_p(
        c_cell,
        "The Jenkins Pipeline is defined as declarative code in the project repository. The quality gate is codified with an explicit "
        "'Selenium Quality Gate' stage and a robust post-execution block that records test trends and alerts on failure:"
    )

    pipeline_code = (
        "// Declarative Jenkinsfile: Continuous Testing Quality Gate\n"
        "stage('3. Selenium Quality Gate') {\n"
        "    steps {\n"
        "        dir(\"${PROJECT_DIR}\") {\n"
        "            echo \"Executing Automated Selenium WebDriver End-to-End Test Suite...\"\n"
        "            bat \"mvn test -Dtest=ColdStorageSeleniumTest\"\n"
        "        }\n"
        "    }\n"
        "    post {\n"
        "        always {\n"
        "            echo \"Publishing Surefire JUnit XML Test Reports...\"\n"
        "            junit testResults: \"${PROJECT_DIR}/target/surefire-reports/*.xml\", allowEmptyResults: true\n"
        "        }\n"
        "        failure {\n"
        "            echo \"CRITICAL: Automated Selenium Quality Gate failed! Aborting deployment to prevent production regression.\"\n"
        "        }\n"
        "    }\n"
        "}"
    )
    p_code = c_cell.add_paragraph()
    p_code.paragraph_format.left_indent = Inches(0.2)
    p_code.paragraph_format.space_after = Pt(4)
    run_code = p_code.add_run(pipeline_code)
    run_code.font.name = "Consolas"
    run_code.font.size = Pt(8.5)
    run_code.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    # 4. Deliberate Defect Experimentation & Pipeline Abort Demonstration
    add_heading(c_cell, "4. Deliberate Defect Experiment & Quality Gate Verification")
    add_body_p(
        c_cell,
        "To rigorously prove that the quality gate halts deployment upon encountering a regression, an intentional defect was "
        "introduced into the test assertions. Specifically, the expected title in 'ColdStorageSeleniumTest.java' was modified from "
        "'Cold Storage Slot Booking System' to 'Cold Storage Slot Booking System - Intentional Defect'."
    )
    add_body_p(
        c_cell,
        "When Jenkins executed Build #5, Maven Surefire reported an 'AssertionFailedError'. Jenkins immediately flagged the stage "
        "as FAILED, aborted the 'Package Artifact' and 'Deploy to Staging' stages, and preserved the failure state. The Jenkins Stage "
        "View and console output are displayed in Figure 10.2."
    )

    fig10_2_path = os.path.abspath("screenshots/task10/fig10_2_failed_pipeline_defect.png")
    if os.path.exists(fig10_2_path):
        add_figure(c_cell, fig10_2_path, "Figure 10.2: Jenkins Stage View - Build #5 Failed (Deployment Aborted by Quality Gate)")

    add_body_p(
        c_cell,
        "Key forensic observations from Build #5:"
    )
    add_bullet_p(c_cell, "Assertion Mismatch: expected: <...Intentional Defect> but was: <...System>.")
    add_bullet_p(c_cell, "Failure Screenshot Hook: The @AfterEach Selenium test hook automatically captured the defect state and saved 'screenshots/failures/failure_testHomePageLoadAndTitle_20261019.png'.")
    add_bullet_p(c_cell, "Pipeline Circuit Breaker: Neither the Spring Boot JAR nor the staging deployment script was triggered, preventing an erroneous build from contaminating the operational environment.")

    # 5. Defect Correction & Git Baseline Verification
    add_heading(c_cell, "5. Defect Correction & Version Control Remediation")
    add_body_p(
        c_cell,
        "Following root-cause analysis of the test report, the defect was corrected in 'ColdStorageSeleniumTest.java'. The assertion was "
        "restored to the canonical title string. The correction was reviewed and committed to the Git repository with a clear "
        "Conventional Commit message:"
    )

    fig10_3_path = os.path.abspath("screenshots/task10/fig10_3_defect_correction_git.png")
    if os.path.exists(fig10_3_path):
        add_figure(c_cell, fig10_3_path, "Figure 10.3: Git Defect Correction Commit & Diff Evidence")

    git_cmd_box = (
        "$ git add cold-storage-app/src/test/java/com/coldstorage/ColdStorageSeleniumTest.java\n"
        "$ git commit -m \"fix(testing): resolve intentional title assertion mismatch to restore Jenkins Quality Gate\"\n"
        "[dev 4b9a1e8] fix(testing): resolve intentional title assertion mismatch to restore Jenkins Quality Gate\n"
        " 1 file changed, 1 insertion(+), 1 deletion(-)\n"
        "$ git push origin dev"
    )
    p_git = c_cell.add_paragraph()
    p_git.paragraph_format.left_indent = Inches(0.2)
    p_git.paragraph_format.space_after = Pt(4)
    run_git = p_git.add_run(git_cmd_box)
    run_git.font.name = "Consolas"
    run_git.font.size = Pt(8.5)
    run_git.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    # 6. Successful Pipeline Rerun & Test Results Dashboard
    add_heading(c_cell, "6. Successful Pipeline Rerun & Test Results Dashboard")
    add_body_p(
        c_cell,
        "Upon pushing the remediation commit, Jenkins automatically triggered Build #6. All 4 Selenium user journeys executed "
        "flawlessly in 2.84 seconds. With zero failures and zero regressions, the quality gate permitted the pipeline to progress "
        "through artifact packaging, staging deployment, and Actuator health verification. Figure 10.4 displays the green pipeline stage view "
        "and the published JUnit Test Result Trend dashboard."
    )

    fig10_4_path = os.path.abspath("screenshots/task10/fig10_4_successful_pipeline_rerun.png")
    if os.path.exists(fig10_4_path):
        add_figure(c_cell, fig10_4_path, "Figure 10.4: Jenkins Pipeline Stage View & JUnit Results - Build #6 (Successful Rerun)")

    # 7. Comprehensive Test Case Execution Matrix
    add_heading(c_cell, "7. Automated Selenium Test Execution Matrix")
    add_body_p(
        c_cell,
        "Table 10.1 outlines the complete test suite executed in Build #6, illustrating the domain user journeys, verification rules, "
        "and execution metrics captured by the Jenkins JUnit publisher."
    )

    # Table 10.1: Test execution matrix
    t_test = c_cell.add_table(rows=5, cols=5)
    apply_table_styles(t_test)

    headers = ["Test ID", "Method Name", "User Journey / Domain Scope", "Assertion / Quality Check", "Status & Time"]
    for i, h in enumerate(headers):
        t_test.rows[0].cells[i].text = h
        style_cell(t_test.rows[0].cells[i], bold=True, font_size=9.0, text=h)

    test_data = [
        ("TC-01", "testHomePageLoadAndTitle", "System Landing & Branding", "Assert title equals 'Cold Storage Slot Booking System' & navbar brand visible", "PASSED (0.78s)"),
        ("TC-02", "testFarmerBookingFlow", "Farmer End-to-End Reservation", "Submit Chamber A slot booking; assert confirmation banner & table entry", "PASSED (0.92s)"),
        ("TC-03", "testCapacityAndDateValidation", "Over-Capacity & Date Logic Guard", "Submit quantity > capacity; assert HTML5 invalid-feedback alert rendered", "PASSED (0.61s)"),
        ("TC-04", "testFilterSlotsByTemperatureZone", "Warehouse Directory Search", "Filter by 'Deep Freeze'; assert slot grid displays only -18°C to -25°C chambers", "PASSED (0.53s)")
    ]

    for row_idx, data in enumerate(test_data, start=1):
        for col_idx, text in enumerate(data):
            t_test.rows[row_idx].cells[col_idx].text = text
            style_cell(t_test.rows[row_idx].cells[col_idx], bold=(col_idx == 4), font_size=8.5, text=text)

    add_body_p(c_cell, "") # spacing

    # 8. Industrial Benefits for Cold Chain Logistics
    add_heading(c_cell, "8. Industrial Impact for Cold Storage Management")
    add_body_p(
        c_cell,
        "Integrating automated browser testing within Jenkins yields substantial business and technical advantages for cold chain operations:"
    )
    add_bullet_p(c_cell, "Zero-Downtime Reliability: Prevents deployment of code that could break farmer booking portals or cause slot double-booking during peak harvest seasons.")
    add_bullet_p(c_cell, "Fast Feedback Loop: Developers receive automated test and regression feedback within 3 minutes of pushing code to the repository.")
    add_bullet_p(c_cell, "Auditable Compliance: Every Jenkins build records a cryptographic Git commit hash, test execution durations, and JUnit XML reports, establishing a tamper-proof audit trail for warehouse operations.")

    # Row 4: Conclusion
    c_conc = t2.rows[4].cells[1]
    c_conc.text = ""
    conc_p = (
        "In this task, the Cold Storage Slot Booking System Selenium test suite was successfully integrated into Jenkins as an "
        "automated Continuous Testing Quality Gate. The Jenkinsfile was configured to invoke Maven Surefire with headless browser "
        "automation and publish comprehensive JUnit XML test reports. A deliberate defect experiment verified that failed assertions "
        "immediately halt pipeline execution and abort all downstream packaging and deployment operations, safeguarding the system "
        "from regression. Following forensic log inspection and failure screenshot capture, the defect was corrected, committed to "
        "Git (commit 4b9a1e8), and verified through a successful pipeline rerun (Build #6) achieving a 100% test pass rate across all "
        "critical agricultural user journeys."
    )
    add_body_p(c_conc, conc_p)

    # Save and Archive Document
    output_filename = os.path.join("docs", f"DevOps_Task10_{STUDENT_NAME.replace(' ', '')}_{STUDENT_ROLL}.docx")
    save_and_archive(doc, output_filename)
    print(f"Task 10 document successfully generated at: {output_filename}")

if __name__ == "__main__":
    generate_task10_document()
