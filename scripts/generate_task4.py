import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from doc_helper import (
    set_student_identity,
    apply_table_styles,
    style_cell,
    add_heading,
    add_subheading,
    add_body_p,
    add_bullet_p,
    add_figure,
    save_and_archive,
    STUDENT_NAME,
    STUDENT_ROLL,
    PROJECT_TITLE
)

def generate_task4_document():
    template_path = os.path.abspath("templates/template_standard.docx")
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template not found at: {template_path}")
        
    doc = docx.Document(template_path)
    
    # Enforce Student Identity in Table 0 and Table 1
    set_student_identity(doc)
    
    # Task 4 uses template_standard.docx which has 5 rows in Table 2:
    # Row 0: Task Number
    # Row 1: Task Title
    # Row 2: Resources / Apparatus Required
    # Row 3: Program & Output (Content)
    # Row 4: Conclusion
    t2 = doc.tables[2]
    
    # Row 0: Task Number
    t2.rows[0].cells[1].text = "4"
    style_cell(t2.rows[0].cells[1], bold=True, font_size=12, text="4")
    
    # Row 1: Task Title
    task_title = (
        "Git and GitHub Repository Initialization: Create a GitHub repository for the Cold Storage "
        "Slot Booking System. Add README, .gitignore, folder structure, issue templates and branch naming rules. "
        "Commit the initial application skeleton using meaningful commit messages."
    )
    t2.rows[1].cells[1].text = task_title
    style_cell(t2.rows[1].cells[1], bold=True, font_size=10.5, text=task_title)

    # Row 2: Resources / Apparatus Required
    hw_text = "Hardware: Workstation / Laptop (x86_64 Intel/AMD), 16 GB DDR4/DDR5 RAM, 512 GB NVMe SSD"
    sw_text = (
        "Software: Git 2.55+, GitHub Enterprise / Cloud Platform, Windows 11 64-bit PowerShell 7, "
        "VS Code / IntelliJ IDEA, Markdown Linter, Git CLI, OpenSSH / HTTPS Credential Helper"
    )
    t2.rows[2].cells[1].text = hw_text
    style_cell(t2.rows[2].cells[1], font_size=9.5, text=hw_text)
    t2.rows[2].cells[2].text = sw_text
    style_cell(t2.rows[2].cells[2], font_size=9.5, text=sw_text)
    
    # Row 3: Program & Output (Main Body Cell)
    body_cell = t2.rows[3].cells[1]
    if body_cell.paragraphs:
        body_cell.paragraphs[0].text = ""
        p0 = body_cell.paragraphs[0]
    else:
        p0 = body_cell.add_paragraph()
        
    p0.paragraph_format.space_before = Pt(4)
    p0.paragraph_format.space_after = Pt(2)
    p0.paragraph_format.keep_with_next = True
    r0 = p0.add_run("1. Version Control Objectives and Repository Architecture")
    r0.bold = True
    r0.font.name = "Calibri"
    r0.font.size = Pt(12)
    
    add_body_p(
        body_cell,
        "The objective of Task 4 is to establish a rock-solid, production-grade Git and GitHub version control foundation for the "
        "Cold Storage Slot Booking System (CSBS). In modern enterprise DevOps pipelines, version control is not merely a backup "
        "mechanism; it serves as the single source of truth that triggers automated continuous integration (CI) builds, orchestrates "
        "Selenium test gates, and powers automated Docker continuous delivery (CD). Without disciplined version control, configuration "
        "drift, broken code merges, and undocumented regressions quickly compromise software quality."
    )
    add_body_p(
        body_cell,
        "To enforce complete engineering discipline, I configured an academic student Git identity (Mayuresh Sawant, 23102A0029@student.edu), "
        "initialized the local repository, crafted an airtight .gitignore file preventing binary bloat, engineered standardized GitHub issue "
        "and pull request templates, and committed the initial Spring Boot application skeleton using Conventional Commits standard."
    )

    # Section 2: Repository Structure & Exclusion Policies
    add_heading(body_cell, "2. Repository Structure & .gitignore Exclusion Standards", keep_with_next=True)
    add_body_p(
        body_cell,
        "To ensure that only source code and version-controlled assets enter Git history, I formulated a strict .gitignore exclusion policy. "
        "Ephemeral artifacts, build outputs (target/), IDE configurations (.idea/, .vscode/), local portable Maven runtimes (maven/), "
        "and archived document drafts (docs/prev/) are excluded to maintain a lean, secure repository. Table 4.1 documents the directory layout:"
    )

    # Table 4.1: Directory Layout (Total = 7900 dxa: 1600, 2400, 3900)
    table_dir = body_cell.add_table(rows=8, cols=3)
    apply_table_styles(table_dir, dxa_widths=[1600, 2400, 3900])
    
    headers_dir = ["Folder / File Path", "Artifact Classification", "DevOps Operational Function & Scope"]
    for i, h in enumerate(headers_dir):
        style_cell(table_dir.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.5, text=h)
        
    dir_data = [
        (".github/", "Workflow & Issue Governance", "Contains bug_report.md, feature_request.md, and pull_request_template.md."),
        ("cold-storage-app/", "Application Source Code", "Complete Maven project containing Spring Boot backend, H2 models, and Thymeleaf UI."),
        ("docs/", "Academic Task Documents", "Single canonical DOCX reports: DevOps_Task{N}_MayureshSawant_23102A0029.docx."),
        ("screenshots/", "Visual Execution Evidence", "Screenshots of UI, terminal outputs, and publication-quality architectural diagrams."),
        ("scripts/", "Automation & Doc Generators", "Python scripts (generate_task{N}.py, doc_helper.py) for 100% reproducible documents."),
        ("templates/", "Department Base Templates", "Approved university layouts (template_task1.docx and template_standard.docx)."),
        (".gitignore", "VCS Exclusion Configuration", "Blocks target/, .class, .log, .idea/, portable maven/, and docs/prev/.")
    ]
    for r_idx, row_vals in enumerate(dir_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_dir.rows[r_idx].cells[c_idx], font_size=8.8, text=val)

    # Section 3: Branching Policy
    add_heading(body_cell, "3. Branch Naming Rules & Branching Strategy", keep_with_next=True)
    add_body_p(
        body_cell,
        "To facilitate seamless multi-developer collaboration, prevent code collisions, and support continuous testing, "
        "I established an enterprise-grade Git branching architecture:"
    )
    add_bullet_p(
        body_cell,
        "Production-ready, protected branch. Direct commits are strictly prohibited. Receives code strictly via tagged releases (e.g., v1.0.0).",
        bold_prefix="main branch: "
    )
    add_bullet_p(
        body_cell,
        "Primary integration branch. Daily feature branch merges undergo automated Jenkins CI build and headless test validation here.",
        bold_prefix="dev branch: "
    )
    add_bullet_p(
        body_cell,
        "Isolated branches branched off 'dev' for individual features (e.g., feature/slot-view, feature/booking-engine). Merged exclusively via PR.",
        bold_prefix="feature/* branches: "
    )
    add_bullet_p(
        body_cell,
        "Dedicated emergency hotfix branches to resolve critical quality gate defects before re-merging into 'dev' and 'main'.",
        bold_prefix="fix/* branches: "
    )

    # Figure 4.1: Branching Diagram
    add_figure(
        body_cell,
        os.path.abspath("screenshots/task4/fig4_1_branching.png"),
        "Figure 4.1: Cold Storage Slot Booking System Git Branching Strategy and Release Flow",
        width_inches=5.4
    )

    # Section 4: GitHub Issue Templates
    add_heading(body_cell, "4. GitHub Issue Templates and Collaboration Contract", keep_with_next=True)
    add_body_p(
        body_cell,
        "To standardize defect reporting and feature enhancement requests across development iterations, I configured standardized "
        "Markdown templates within .github/ISSUE_TEMPLATE/: Table 4.2 outlines the issue tracking policies:"
    )

    # Table 4.2: Issue & PR Templates (Total = 7900 dxa: 1500, 1800, 4600)
    table_tmpl = body_cell.add_table(rows=4, cols=3)
    apply_table_styles(table_tmpl, dxa_widths=[1500, 1800, 4600])
    
    headers_tmpl = ["Template Name", "Target File Path", "Enforced Fields & Quality Checkpoints"]
    for i, h in enumerate(headers_tmpl):
        style_cell(table_tmpl.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.5, text=h)
        
    tmpl_data = [
        ("Bug Report", ".github/ISSUE_TEMPLATE/bug_report.md", 
         "Defect summary, affected module checklist, step-by-step reproduction, expected vs actual behavior, and environment metadata."),
        ("Feature Request", ".github/ISSUE_TEMPLATE/feature_request.md", 
         "User story problem statement, proposed technical solution, target cold storage flow, and architectural context."),
        ("Pull Request", ".github/pull_request_template.md", 
         "Summary of changes, linked issue references, branch category, and mandatory 5-point quality and zero-regression checklist.")
    ]
    for r_idx, row_vals in enumerate(tmpl_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_tmpl.rows[r_idx].cells[c_idx], font_size=8.8, text=val)

    # Section 5: Initial Application Skeleton Commit Log
    add_heading(body_cell, "5. Initial Application Skeleton Commit Log and Graph", keep_with_next=True)
    add_body_p(
        body_cell,
        "The repository was initialized and populated with two foundational commits following Conventional Commits format. "
        "The commit graph and branch states were inspected in terminal and are summarized in Table 4.3:"
    )

    # Table 4.3: Initial Commits (Total = 7900 dxa: 1000, 1200, 3500, 2200)
    table_cmt = body_cell.add_table(rows=3, cols=4)
    apply_table_styles(table_cmt, dxa_widths=[1000, 1200, 3500, 2200])
    
    headers_cmt = ["Commit Hash", "Branch", "Conventional Commit Message", "Committed Assets & Impact"]
    for i, h in enumerate(headers_cmt):
        style_cell(table_cmt.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.5, text=h)
        
    cmt_data = [
        ("17be392", "main", "feat(arch): initialize Cold Storage Slot Booking System skeleton and Maven architecture",
         "61 files, 4197 insertions. Baselined Spring Boot app, models, UI templates, and initial tasks."),
        ("d014f24", "dev (HEAD)", "docs: establish dev branch baseline and issue templates",
         "Created dev integration branch, configured branch naming rules, and established PR template.")
    ]
    for r_idx, row_vals in enumerate(cmt_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_cmt.rows[r_idx].cells[c_idx], font_size=8.8, text=val)

    # Figure 4.2: Terminal Output Visual
    add_figure(
        body_cell,
        os.path.abspath("screenshots/task4/fig4_2_git_terminal.png"),
        "Figure 4.2: Terminal Git Log Graph, Branch Topology, and Student Identity Verification",
        width_inches=5.4
    )

    # Row 4: Conclusion
    conclusion_text = (
        "In this task, I successfully established the official Git and GitHub version control repository for the "
        "Cold Storage Slot Booking System (CSBS). I configured the student author identity (Mayuresh Sawant, 23102A0029), "
        "crafted an airtight .gitignore exclusion policy protecting the repository against build and runtime clutter, and engineered "
        "standardized issue and pull request templates within .github/. I codified an enterprise Git branching strategy centered "
        "on protected 'main' and integration 'dev' branches, complemented by structured 'feature/*' branches following Conventional "
        "Commit conventions. The initial application skeleton was committed with complete fidelity, establishing a pristine, "
        "traceable source baseline ready for feature branching in Task 5 and collaborative MVP delivery in Task 6."
    )
    t2.rows[4].cells[1].text = conclusion_text
    style_cell(t2.rows[4].cells[1], font_size=10.5, text=conclusion_text)
    
    output_path = os.path.abspath(f"docs/DevOps_Task4_{STUDENT_NAME.replace(' ', '')}_{STUDENT_ROLL}.docx")
    saved_file = save_and_archive(doc, output_path)
    print(f"Task 4 document successfully generated at: {saved_file}")

if __name__ == "__main__":
    generate_task4_document()
