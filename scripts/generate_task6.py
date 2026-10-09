import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from doc_helper import (
    set_student_identity,
    apply_table_styles,
    style_cell,
    add_heading,
    add_subheading,
    add_body_p,
    add_bullet_p,
    add_figure,
    save_and_archive,
    STUDENT_NAME,
    STUDENT_ROLL,
    PROJECT_TITLE
)

def generate_task6_document():
    template_path = os.path.abspath("templates/template_standard.docx")
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template not found at: {template_path}")
        
    doc = docx.Document(template_path)
    
    # Enforce Student Identity in Table 0 and Table 1
    set_student_identity(doc)
    
    # Task 6 uses template_standard.docx which has 5 rows in Table 2:
    # Row 0: Task Number
    # Row 1: Task Title
    # Row 2: Resources / Apparatus Required
    # Row 3: Program & Output (Content)
    # Row 4: Conclusion
    t2 = doc.tables[2]
    
    # Row 0: Task Number
    t2.rows[0].cells[1].text = "6"
    style_cell(t2.rows[0].cells[1], bold=True, font_size=12, text="6")
    
    # Row 1: Task Title
    task_title = (
        "MVP Completion and Git Collaboration: Implement the remaining essential MVP functions for "
        "the Cold Storage Slot Booking System: user registration, availability/slot view, booking request, "
        "confirmation, cancellation and status tracking. Demonstrate a second feature branch, "
        "conflict creation/resolution, tagging and a release-ready source baseline."
    )
    t2.rows[1].cells[1].text = task_title
    style_cell(t2.rows[1].cells[1], bold=True, font_size=10.5, text=task_title)

    # Row 2: Resources / Apparatus Required
    hw_text = "Hardware: Workstation / Laptop (x86_64 Intel/AMD), 16 GB DDR4/DDR5 RAM, 512 GB NVMe SSD"
    sw_text = (
        "Software: Git 2.55+, GitHub Enterprise / Cloud Platform, Windows 11 64-bit PowerShell 7, "
        "Java 21 LTS, Spring Boot 3.3.4, Apache Maven 3.9.9, Embedded H2 Database, VS Code / IntelliJ IDEA"
    )
    t2.rows[2].cells[1].text = hw_text
    style_cell(t2.rows[2].cells[1], font_size=9.5, text=hw_text)
    t2.rows[2].cells[2].text = sw_text
    style_cell(t2.rows[2].cells[2], font_size=9.5, text=sw_text)
    
    # Row 3: Program & Output (Main Body Cell)
    body_cell = t2.rows[3].cells[1]
    if body_cell.paragraphs:
        body_cell.paragraphs[0].text = ""
        p0 = body_cell.paragraphs[0]
    else:
        p0 = body_cell.add_paragraph()
        
    p0.paragraph_format.space_before = Pt(4)
    p0.paragraph_format.space_after = Pt(2)
    p0.paragraph_format.keep_with_next = True
    r0 = p0.add_run("1. MVP Completion and Collaborative Engineering Objectives")
    r0.bold = True
    r0.font.name = "Calibri"
    r0.font.size = Pt(12)
    
    add_body_p(
        body_cell,
        "The objective of Task 6 is twofold: first, to deliver the fully functional Minimum Viable Product (MVP) of the "
        "Cold Storage Slot Booking System (CSBS) covering all five core operational modules; and second, to demonstrate advanced Git "
        "collaboration techniques by engineering a controlled merge conflict on a shared service component, analyzing Git conflict markers, "
        "manually harmonizing conflicting requirements, and establishing an official release-ready source baseline tagged 'v1.0.0'."
    )
    add_body_p(
        body_cell,
        "In production DevOps environments with concurrent developer teams, merge conflicts are an inevitable reality of parallel feature development. "
        "Understanding how Git identifies collisions through the 3-way merge algorithm and mastering manual resolution guarantees that "
        "unintended regressions, lost logic, or syntactical breakage do not corrupt the integration branch prior to CI build triggers."
    )

    # Section 2: Summary of Functional MVP Modules
    add_heading(body_cell, "2. Summary of Fully Functional MVP Modules", keep_with_next=True)
    add_body_p(
        body_cell,
        "With the completion of Sprint 2, all five core MVP modules are fully implemented, integrated, and verified on http://localhost:8080:"
    )
    add_bullet_p(
        body_cell,
        "Implemented in AuthController, UserService, and UserRepository. Provides secure registration, role assignment (Farmer, Trader, Manager), and authentication.",
        bold_prefix="Module 1: User & Farmer Registration/Auth: "
    )
    add_bullet_p(
        body_cell,
        "Implemented in SlotController, ChamberService, and slots.html. Visualizes chambers across Deep Freeze (-25°C to -18°C), Chilled (0°C to 4°C), and CA zones.",
        bold_prefix="Module 2: Chamber & Slot Availability Directory: "
    )
    add_bullet_p(
        body_cell,
        "Implemented in BookingController, BookingService, and booking-form.html. Captures produce type, crate quantity, and date windows with automated cost computation.",
        bold_prefix="Module 3: Slot Booking & Reservation Engine: "
    )
    add_bullet_p(
        body_cell,
        "Implemented in BookingService.findOverlappingBookings(). Rejects double-booking collisions on the same chamber slot during overlapping date intervals.",
        bold_prefix="Module 4: Date Overlap & Capacity Validation Guard: "
    )
    add_bullet_p(
        body_cell,
        "Implemented in BookingController lifecycle actions. Governs state transitions: PENDING -> CONFIRMED -> CHECKED_IN -> COMPLETED / CANCELLED.",
        bold_prefix="Module 5: Lifecycle Management & Status Machine: "
    )

    # Section 3: Merge Conflict Experiment
    add_heading(body_cell, "3. Intentional Merge Conflict Creation on 'feature/booking-rules'", keep_with_next=True)
    add_body_p(
        body_cell,
        "To rigorously demonstrate conflict resolution, I spawned a second feature branch named 'feature/booking-rules' off 'dev'. "
        "Both branches simultaneously modified the booking reference formatting logic within cold-storage-app/.../BookingService.java:"
    )
    add_bullet_p(
        body_cell,
        "Modified line 76 to generate booking references prefixed with 'CSBS-RES-' (e.g., CSBS-RES-A1B2C3D4), committing changes in commit 1d657e3.",
        bold_prefix="On feature/booking-rules: "
    )
    add_bullet_p(
        body_cell,
        "Modified the exact same line 76 to enforce an academic year prefix 'CSBS-2026-' (e.g., CSBS-2026-A1B2C3D4), committing changes in commit a5b4c22.",
        bold_prefix="On dev integration branch: "
    )

    # Section 4: Conflict Analysis & Manual Harmonization
    add_heading(body_cell, "4. Conflict Inspection, Marker Analysis & Harmonization", keep_with_next=True)
    add_body_p(
        body_cell,
        "When executing 'git merge feature/booking-rules' into 'dev', Git's merge engine detected concurrent textual modifications "
        "on the identical line block and aborted automatic integration with a CONFLICT error. Git inserted standardized conflict markers:"
    )
    add_bullet_p(
        body_cell,
        "Signifies the conflicting content existing on the current target branch ('dev').",
        bold_prefix="<<<<<<< HEAD: "
    )
    add_bullet_p(
        body_cell,
        "Separates the target branch code from the incoming branch code.",
        bold_prefix="=======: "
    )
    add_bullet_p(
        body_cell,
        "Signifies the incoming conflicting code arriving from 'feature/booking-rules'.",
        bold_prefix=">>>>>>> feature/booking-rules: "
    )
    add_body_p(
        body_cell,
        "I manually resolved the conflict by harmonizing both business requirements: combining the academic year identifier and the reservation "
        "indicator into a unified standard: 'CSBS-2026-RES-' + UUID. I staged the resolved file and completed the merge via commit 15b1c98."
    )

    # Figure 6.1: Conflict Diagram
    add_figure(
        body_cell,
        os.path.abspath("screenshots/task6/fig6_1_conflict_workflow.png"),
        "Figure 6.1: Git Collaboration, Intentional Merge Conflict Creation, and Resolution Architecture",
        width_inches=5.4
    )

    # Figure 6.2: Conflict Terminal Visual
    add_figure(
        body_cell,
        os.path.abspath("screenshots/task6/fig6_2_conflict_terminal.png"),
        "Figure 6.2: Terminal Output Demonstrating Conflict Markers and Manual Code Harmonization",
        width_inches=5.4
    )

    # Section 5: Semantic Version Tagging
    add_heading(body_cell, "5. Semantic Version Tagging (v1.0.0) & Production Merge", keep_with_next=True)
    add_body_p(
        body_cell,
        "Following conflict resolution, the stabilized MVP baseline on 'dev' was tagged with an annotated semantic version: "
        "'git tag -a v1.0.0 -m \"Release v1.0.0: Fully Functional Cold Storage Slot Booking System MVP\"'. "
        "The verified release was subsequently merged into 'main' via a non-fast-forward merge commit (60ffbc1), establishing "
        "a clean, release-ready production baseline. Table 6.1 summarizes the complete Git commit tree:"
    )

    # Table 6.1: Commit History (Total = 7900 dxa: 1000, 1400, 3500, 2000)
    table_cmt = body_cell.add_table(rows=7, cols=4)
    apply_table_styles(table_cmt, dxa_widths=[1000, 1400, 3500, 2000])
    
    headers_cmt = ["Commit Hash", "Branch / Tag", "Conventional Commit Message", "Engineering Milestones"]
    for i, h in enumerate(headers_cmt):
        style_cell(table_cmt.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.5, text=h)
        
    cmt_data = [
        ("60ffbc1", "main", "release: merge dev baseline into main for production release v1.0.0", "Production release merge from dev into main."),
        ("15b1c98", "tag: v1.0.0, dev", "fix(merge): resolve booking reference format conflict", "Resolved conflict; harmonized CSBS-2026-RES- reference."),
        ("1d657e3", "feature/booking-rules", "feat(booking): implement CSBS-RES reference prefix", "Conflicting change committed on feature branch."),
        ("a5b4c22", "dev", "feat(booking): set CSBS-2026 reference format on dev branch", "Conflicting change committed on integration branch."),
        ("3e45a1f", "dev", "merge: pull request #1 from feature/slot-view into dev", "Integrated Feature 1 (Chamber Directory)."),
        ("3756f4e", "feature/slot-view", "feat(slot): enhance chamber directory with commodity search", "Feature 1 service & UI implementation.")
    ]
    for r_idx, row_vals in enumerate(cmt_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_cmt.rows[r_idx].cells[c_idx], font_size=8.5, text=val)

    # Figure 6.3: Tag and Release Tree
    add_figure(
        body_cell,
        os.path.abspath("screenshots/task6/fig6_3_tag_release_tree.png"),
        "Figure 6.3: Terminal Git Tag Verification (v1.0.0) and Complete Branch Commit Graph",
        width_inches=5.4
    )

    # Section 6: Updated Agile Backlog
    add_heading(body_cell, "6. Updated Product Backlog (Sprint 2 Milestone Completion)", keep_with_next=True)
    add_body_p(
        body_cell,
        "With the completion of Sprint 2, all functional MVP user stories have transitioned to DONE. "
        "Table 6.2 records the updated status of product backlog items:"
    )

    # Table 6.2: Updated Backlog (Total = 7900 dxa: 900, 2500, 1000, 1200, 1100, 1200)
    table_ub = body_cell.add_table(rows=8, cols=6)
    apply_table_styles(table_ub, dxa_widths=[900, 2500, 1000, 1200, 1100, 1200])
    
    headers_ub = ["Item ID", "User Story / Feature Description", "Story Points", "Sprint", "DoD Status", "Verification"]
    for i, h in enumerate(headers_ub):
        style_cell(table_ub.rows[0].cells[i], bg_color="EAECEE", bold=True, font_size=9.0, text=h)
        
    ub_data = [
        ("PB-01", "Application Foundation & H2 Data Models", "5 SP", "Sprint 1", "DONE", "Tested in Task 3"),
        ("PB-02", "Git VCS Setup & Branching Architecture", "3 SP", "Sprint 1", "DONE", "Verified in Task 4"),
        ("PB-03", "Chamber Directory & Slot Grid UI", "5 SP", "Sprint 2", "DONE", "Merged in Task 5"),
        ("PB-04", "Slot Booking & Date Overlap Validator", "5 SP", "Sprint 2", "DONE", "Verified in Task 6"),
        ("PB-05", "Manager Lifecycle Status Machine", "3 SP", "Sprint 2", "DONE", "Tested in Task 6"),
        ("PB-06", "Jenkins CI Build & Artifact Archiving", "8 SP", "Sprint 3", "READY", "Target for Task 7"),
        ("PB-07", "Selenium Automated Quality Gate Suite", "8 SP", "Sprint 3", "READY", "Target for Task 9")
    ]
    for r_idx, row_vals in enumerate(ub_data, start=1):
        for c_idx, val in enumerate(row_vals):
            style_cell(table_ub.rows[r_idx].cells[c_idx], font_size=8.5, text=val)

    # Figure 6.4: Working MVP Screen
    add_figure(
        body_cell,
        os.path.abspath("screenshots/task6/fig6_4_mvp_booking_live.png"),
        "Figure 6.4: Verified Working MVP: Real-Time Bookings Directory and Lifecycle Action Controls",
        width_inches=5.4
    )

    # Row 4: Conclusion
    conclusion_text = (
        "In this task, I successfully completed all essential MVP functions of the Cold Storage Slot Booking System (CSBS) "
        "and demonstrated production-grade Git collaboration and release management. I implemented and verified user registration, "
        "multi-temperature chamber browsing, slot reservations with date-overlap validation, and role-based status lifecycle actions. "
        "To validate team collaboration workflows, I created an intentional merge conflict on BookingService.java using a second feature branch "
        "('feature/booking-rules'), analyzed the raw conflict markers, and manually harmonized conflicting reference prefix policies into a clean "
        "resolution commit (15b1c98). I tagged the verified baseline as 'v1.0.0', merged the release into 'main', and transitioned all Sprint 2 "
        "backlog items to DONE. This establishes a pristine, tested source baseline ready for Jenkins Continuous Integration in Task 7."
    )
    t2.rows[4].cells[1].text = conclusion_text
    style_cell(t2.rows[4].cells[1], font_size=10.5, text=conclusion_text)
    
    output_path = os.path.abspath(f"docs/DevOps_Task6_{STUDENT_NAME.replace(' ', '')}_{STUDENT_ROLL}.docx")
    saved_file = save_and_archive(doc, output_path)
    print(f"Task 6 document successfully generated at: {saved_file}")

if __name__ == "__main__":
    generate_task6_document()
